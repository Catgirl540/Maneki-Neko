# -*- coding: utf-8 -*-
"""检查 DOCX 中是否残留 Markdown 标记，并统计结构。"""
import os
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
doc = Document(os.path.join(HERE, "招财猫银行-技术文档.docx"))

BAD = ("**", "`")

def scan(text):
    return [m for m in BAD if m in text]

bad = []
for p in doc.paragraphs:
    if scan(p.text):
        bad.append(("P", p.text[:70]))
for tb in doc.tables:
    for row in tb.rows:
        for cell in row.cells:
            if scan(cell.text):
                bad.append(("T", cell.text[:70]))

print("leftover count:", len(bad))
for kind, txt in bad[:10]:
    print(" ", kind, repr(txt))

print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables),
      "images:", len(doc.inline_shapes))
print("headings:", sum(1 for p in doc.paragraphs if p.style.name.startswith("Heading")))
